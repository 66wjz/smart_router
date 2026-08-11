"""读取 AI中转站模型可用清单v2.xlsx，输出模型参数表"""
import openpyxl

path = r"..\docs\AI中转站模型可用清单v2.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
for ws in wb.worksheets:
    print(f"\n===== Sheet: {ws.title}  ({ws.max_row} 行 x {ws.max_column} 列) =====")
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i > 60:
            print("...（截断）")
            break
        cells = [str(c).strip() if c is not None else "" for c in row]
        if any(cells):
            print(f"  {i}: " + " | ".join(cells))
